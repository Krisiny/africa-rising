"""Create content.xlsx with the original game content.

    python make_content.py            create content.xlsx (never overwrites it)
    python make_content.py --force    replace content.xlsx (erases edits made in Excel!)
    python make_content.py --out FILE write to another file instead

After this, edit the texts and numbers in Excel. The game reads content.xlsx,
not this script.
"""

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

import config

ROOT = Path(__file__).resolve().parent

COUNTRY_HEADERS = ["code", "name", "color", "card"]
COUNTRIES = [
    ("NG", "Nigeria", "#0E8A4F", "Africa's most populous country; Oil and gas exporter; Lagos tech hub; Afrobeats and Nollywood"),
    ("EG", "Egypt", "#C8930A", "Suez Canal; Pyramids and Red Sea tourism; Nearly all its water comes from the Nile; Big wheat importer"),
    ("KE", "Kenya", "#B3202A", "M-Pesa mobile money and tech hub; Safaris; Tea, flowers and avocados; About 90% renewable power"),
    ("ZA", "South Africa", "#6B3FA0", "Africa's most industrialized economy; Mining (platinum, gold); Car factories; Very high unemployment"),
    ("CD", "DR Congo", "#2E86C1", "About 70% of the world's cobalt; Congo River with huge hydropower potential; Rainforest; Conflict in the east"),
]

EVENT_HEADERS = ["id", "title", "kind", "description", "option_1", "option_2", "option_3"]
EVENTS = [
    (1, "Foreign investor", "good", "A foreign investor offers $1 billion. Where should it go?",
     "Oil & mining", "Tourism", "Tech & startups"),
    (2, "Blackouts", "bad", "Power cuts are shutting down factories and homes. What do you build?",
     "Gas power plant", "Solar & wind farms", "Hydro dam"),
    (3, "Tourism boom", "good", "A travel video about Africa goes viral and millions want to visit. What do you promote?",
     "Wildlife & nature", "History & culture", "Beach resorts"),
    (4, "Drought", "bad", "The rains failed. Crops are dying and food prices are rising. How do you respond?",
     "Import food", "Irrigation & water saving", "Cash to farmers by phone"),
    (5, "Electric car boom", "good", "The world is switching to electric cars and needs batteries. How do you cash in?",
     "Mine battery metals", "Build car factories", "Electric motorbikes & buses"),
    (6, "Youth unemployment", "bad", "Millions of young people can't find work and protests are starting. Where do you create jobs?",
     "Tech & digital jobs", "New factories", "Modern farming"),
    (7, "Africa Cup of Nations", "good", "You can host the Africa Cup of Nations, Africa's biggest football tournament. What do you do?",
     "Host it alone", "Co-host with neighbors", "Skip it, build roads"),
    (8, "Shipping crisis", "bad", "Attacks on ships in the Red Sea push trade away from the Suez Canal and around Africa. How do you react?",
     "Upgrade your ports", "Cut fees for ships", "Trade more with neighbors"),
]

RESULT_HEADERS = ["event_id", "country", "option_1", "option_2", "option_3", "reason"]
RESULTS = [
    (1, "NG", 2, 0, 3, "Oil dominates Nigeria's exports but creates few jobs, while Lagos is one of Africa's biggest tech hubs."),
    (1, "EG", 1, 3, 1, "Tourism is one of Egypt's biggest earners, from the pyramids to Red Sea beach resorts."),
    (1, "KE", 0, 2, 3, "Kenya is nicknamed the Silicon Savannah: it's the home of M-Pesa and a leading tech hub."),
    (1, "ZA", 3, 1, 2, "South Africa is the world's top platinum producer, and mining is a pillar of its exports."),
    (1, "CD", 3, 0, 0, "DR Congo mines about 70% of the world's cobalt, a key metal in phone and car batteries."),
    (2, "NG", 1, 0, -1, "Nigeria has Africa's largest gas reserves, and most of its electricity already comes from gas."),
    (2, "EG", -1, 1, -2, "Egypt's desert is ideal for solar, and gas shortages were behind its 2024 blackouts."),
    (2, "KE", -2, 1, -1, "Kenya's power is already about 90% renewable; its dams struggle in droughts, so sun and wind are safer."),
    (2, "ZA", -1, 1, -2, "South Africa has plenty of sun and wind but little gas and few big rivers, and solar farms go up fast."),
    (2, "CD", -2, 0, 1, "The Congo River's Inga Falls has one of the biggest hydropower potentials on Earth."),
    (3, "NG", 0, 3, 1, "Afrobeats, Nollywood and Lagos' festive 'Detty December' concert season pull in huge crowds."),
    (3, "EG", 0, 3, 2, "The pyramids are Egypt's world-famous brand, and its Red Sea beaches come a close second."),
    (3, "KE", 3, 1, 1, "Safaris in the Maasai Mara, where the Great Migration arrives every year, are Kenya's tourism magnet."),
    (3, "ZA", 3, 1, 2, "Kruger National Park and Table Mountain draw visitors from around the world."),
    (3, "CD", 1, 0, -1, "Virunga's mountain gorillas are world-famous, though conflict in the east keeps many tourists away."),
    (4, "NG", -1, 1, 0, "A weak naira makes imports expensive, so irrigating northern farms grows food at home."),
    (4, "EG", 1, 0, -1, "Egypt farms with Nile water, not rain, and is already one of the world's biggest wheat importers."),
    (4, "KE", 0, 0, 1, "Most Kenyan adults use M-Pesa, so cash aid reaches farmers' phones in minutes."),
    (4, "ZA", 0, 1, 0, "In 2018, strict water-saving rules saved Cape Town from running dry on 'Day Zero'."),
    (4, "CD", -1, 1, 0, "DR Congo is full of rivers but irrigates almost none of its farmland."),
    (5, "NG", 2, 1, 1, "Nigeria has lithium, a key battery metal, and Chinese firms are building plants to process it there."),
    (5, "EG", 0, 2, 1, "Egypt's Suez Canal Economic Zone links Europe, Asia and Africa and keeps attracting factories."),
    (5, "KE", -1, 0, 3, "Kenya's power is mostly renewable and its roads are full of motorbike taxis, perfect for e-bikes."),
    (5, "ZA", 2, 3, 0, "South Africa already builds cars for BMW, Mercedes, VW and Toyota and exports them to Europe."),
    (5, "CD", 3, -1, 0, "DR Congo has most of the world's cobalt and is the world's second-biggest copper producer."),
    (6, "NG", 1, 0, 0, "Young Nigerians already power booming industries like fintech, Nollywood and Afrobeats."),
    (6, "EG", 0, 1, -1, "Egypt has a huge workforce and growing factory zones near the Suez Canal."),
    (6, "KE", 0, -1, 1, "Farming is Kenya's biggest employer, and tea, flowers and avocados are top exports."),
    (6, "ZA", 1, 0, 0, "Call centers serving the UK, US and Australia are one of South Africa's fastest-growing youth employers."),
    (6, "CD", -1, 0, 1, "DR Congo has one of the largest areas of unused farmland in the world."),
    (7, "NG", 0, 2, 1, "Nigeria co-hosted the 2000 tournament with Ghana, and sharing cuts the stadium bill."),
    (7, "EG", 2, 1, 1, "Egypt already has big modern stadiums and hosted the 2019 tournament with only months of notice."),
    (7, "KE", -1, 2, 1, "Kenya won the right to co-host the 2027 tournament with Uganda and Tanzania, sharing the cost."),
    (7, "ZA", 2, 1, 1, "South Africa can reuse the stadiums it built for the 2010 World Cup."),
    (7, "CD", -2, -1, 2, "DR Congo has very few paved roads for a country the size of Western Europe, so roads pay off more."),
    (8, "NG", 1, 0, 0, "Lagos opened the deep-sea Lekki port in 2023, ready for big ships rerouted past West Africa."),
    (8, "EG", -1, -1, 0, "Ships avoided the canal for safety, not price, so Egypt's canal income fell by more than half in 2024."),
    (8, "KE", 0, -1, 1, "Mombasa is the main port for landlocked neighbors like Uganda and South Sudan, so regional trade softens the blow."),
    (8, "ZA", 1, 0, 0, "Ships now sail around the Cape, but South Africa's ports rank among the world's slowest."),
    (8, "CD", -1, 0, 1, "DR Congo has almost no coastline and exports its metals through neighbors' ports."),
]

# Column widths (in Excel characters) and which columns wrap their text.
WIDTHS = {
    "countries": [8, 16, 11, 90],
    "events": [6, 24, 8, 60, 28, 28, 28],
    "results": [10, 10, 11, 11, 11, 90],
}
WRAP = {
    "countries": {"card"},
    "events": {"description", "option_1", "option_2", "option_3"},
    "results": {"reason"},
}
SIGNED = "+0;-0;0"   # Excel shows +2, -1 and 0


def add_sheet(wb, name, headers, rows):
    """Add one sheet: bold frozen header row, then the data rows."""
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, (header, width) in enumerate(zip(headers, WIDTHS[name]), start=1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = width
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.alignment = Alignment(vertical="top", wrap_text=header in WRAP[name])
            if name == "results" and header.startswith("option_"):
                cell.number_format = SIGNED


def build() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)  # drop the default empty sheet
    add_sheet(wb, "countries", COUNTRY_HEADERS, COUNTRIES)
    add_sheet(wb, "events", EVENT_HEADERS, EVENTS)
    add_sheet(wb, "results", RESULT_HEADERS, RESULTS)
    return wb


def main() -> int:
    parser = argparse.ArgumentParser(description="Create content.xlsx with the original game content.")
    parser.add_argument("--force", action="store_true", help="replace the file if it already exists")
    parser.add_argument("--out", type=Path, help="where to write (default: the game's content.xlsx)")
    args = parser.parse_args()

    out = args.out or ROOT / config.CONTENT_FILE
    if out.exists() and not args.force:
        print(f"{out.name} already exists, so nothing was written.")
        print("Run python make_content.py --force to replace it (this erases any edits made in Excel).")
        return 1
    try:
        build().save(out)
    except OSError as err:
        print(f"Could not write {out.name} ({type(err).__name__}). If it is open in Excel, close it and try again.")
        return 1
    print(f"Created {out.name}: {len(COUNTRIES)} countries, {len(EVENTS)} events, {len(RESULTS)} results.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
