"""Load and validate content.xlsx (countries, events and results).

No pygame here, so this can be used by `python main.py --check`.
"""

from __future__ import annotations

import io
import json
import re
import warnings
from dataclasses import dataclass, field
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:   # the web version reads content.json instead (see make_web.py)
    load_workbook = get_column_letter = None


@dataclass(frozen=True)
class Country:
    code: str                      # "KE"
    name: str                      # "Kenya"
    color: tuple[int, int, int]    # team color as RGB
    card: tuple[str, ...]          # 3 or 4 short facts, one per line


@dataclass(frozen=True)
class Event:
    id: str                        # event id as text, e.g. "1"
    title: str
    kind: str                      # "good" or "bad"
    description: str
    options: tuple[str, str, str]


@dataclass(frozen=True)
class Result:
    values: tuple[int, int, int]   # GDP change in percentage points per option
    reason: str                    # why the best option is best for this country

    @property
    def best_value(self) -> int:
        return max(self.values)

    @property
    def best_indices(self) -> tuple[int, ...]:
        """0-based indices of every option that ties for the highest value."""
        return tuple(i for i, v in enumerate(self.values) if v == self.best_value)


@dataclass
class Content:
    countries: list[Country]                  # turn order = row order
    events: list[Event]                       # board order = row order
    results: dict[tuple[str, str], Result] = field(default_factory=dict)  # (event_id, code)

    @property
    def n_tiles(self) -> int:
        return len(self.events)

    @property
    def codes(self) -> list[str]:
        return [c.code for c in self.countries]

    def country(self, code: str) -> Country:
        for c in self.countries:
            if c.code == code:
                return c
        raise KeyError(code)

    def event_at_tile(self, tile: int) -> Event:
        """Tiles are numbered from 1."""
        return self.events[tile - 1]

    def event_by_id(self, event_id: str) -> Event:
        for e in self.events:
            if e.id == event_id:
                return e
        raise KeyError(event_id)

    def result(self, event_id: str, code: str) -> Result:
        return self.results[(event_id, code)]

    def to_json(self) -> dict:
        """Plain data for content.json (used by the web version, which has no Excel reader)."""
        return {
            "countries": [{"code": c.code, "name": c.name, "color": list(c.color), "card": list(c.card)}
                          for c in self.countries],
            "events": [{"id": e.id, "title": e.title, "kind": e.kind, "description": e.description,
                        "options": list(e.options)} for e in self.events],
            "results": [{"event_id": eid, "country": code, "values": list(r.values), "reason": r.reason}
                        for (eid, code), r in self.results.items()],
        }

    @classmethod
    def from_json(cls, data: dict) -> Content:
        countries = [Country(c["code"], c["name"], tuple(c["color"]), tuple(c["card"]))
                     for c in data["countries"]]
        events = [Event(e["id"], e["title"], e["kind"], e["description"], tuple(e["options"]))
                  for e in data["events"]]
        results = {(r["event_id"], r["country"]): Result(tuple(r["values"]), r["reason"])
                   for r in data["results"]}
        return cls(countries, events, results)


@dataclass(frozen=True)
class Problem:
    """One thing that is wrong in content.xlsx, written for the people editing it."""
    sheet: str | None              # "results", or None for file-level problems
    row: int | None                # Excel row number (header = row 1), or None
    column: str | None             # e.g. "D (option_1)", or None
    message: str                   # what is wrong
    fix: str                       # how to fix it

    def where(self) -> str:
        parts = []
        if self.sheet:
            parts.append(f"Sheet '{self.sheet}'")
        if self.row is not None:
            parts.append(f"row {self.row}")
        if self.column:
            parts.append(f"column {self.column}")
        return ", ".join(parts)

    def __str__(self) -> str:
        where = self.where()
        prefix = f"{where}: " if where else ""
        return f"{prefix}{self.message} Fix: {self.fix}"


class ContentError(Exception):
    """Raised by load_content with every problem found, not just the first."""

    def __init__(self, problems: list[Problem]):
        self.problems = problems
        super().__init__(f"{len(problems)} problem(s) in the content file")


# --- Reading the workbook ----------------------------------------------------

# Sheets and their columns, in the order they appear in content.xlsx.
HEADERS = {
    "countries": ("code", "name", "color", "card"),
    "events": ("id", "title", "kind", "description", "option_1", "option_2", "option_3"),
    "results": ("event_id", "country", "option_1", "option_2", "option_3", "reason"),
}
OPTIONS = ("option_1", "option_2", "option_3")
COUNTRY_LIMITS = (2, 6)
EVENT_LIMITS = (4, 12)      # what the board layout fits
FACT_LIMITS = (3, 4)

# What belongs in each column, for the "how to fix it" part of a problem.
WHAT = {
    "code": "a short country code such as KE",
    "name": "the country name",
    "color": "a team color such as #B3202A",
    "card": "3 or 4 short facts separated by semicolons (;)",
    "id": "an event id such as 1",
    "title": "a short event title",
    "kind": "good or bad",
    "description": "a sentence or two about the event",
    "option_1": "the text of option 1",
    "option_2": "the text of option 2",
    "option_3": "the text of option 3",
    "event_id": "an event id from the events sheet",
    "country": "a country code from the countries sheet",
    "reason": "one sentence on why the best option is best for this country",
}
NUMBER_FIX = "Type the GDP change as a whole number such as +2, 0 or -1 (no decimals, words or % sign)."


def load_content(path: Path) -> Content:
    """Read and validate the content file. Raises ContentError listing all problems."""
    path = Path(path)
    if path.suffix.lower() == ".json":      # web version: content.xlsx already checked and baked in
        try:
            return Content.from_json(json.loads(path.read_text(encoding="utf-8")))
        except Exception as err:
            raise ContentError([Problem(None, None, None, f"{path.name} could not be read ({err}).",
                                        "Build the web version again with python make_web.py.")])
    problems: list[Problem] = []
    workbook = _open_workbook(path, problems)
    if workbook is None:
        raise ContentError(problems)
    content = None
    try:
        content = _read_content(workbook, problems)
    except Exception as err:  # last-resort safety net: never a traceback in class
        problems.append(Problem(
            None, None, None,
            f"Something unexpected went wrong while reading {path.name} ({type(err).__name__}: {err}).",
            "Open the file in Excel, check it, and save it again as an Excel Workbook (.xlsx)."))
    finally:
        workbook.close()
    if problems:
        raise ContentError(problems)
    return content


def _open_workbook(path: Path, problems: list[Problem]):
    """Return the openpyxl workbook, or None after adding a file-level problem."""
    name = path.name
    if not path.is_file():
        problems.append(Problem(
            None, None, None,
            f"The file {name} was not found in the folder {path.parent}.",
            f"Put {name} in the game folder next to main.py (or set CONTENT_FILE in config.py)."))
        return None
    try:
        data = path.read_bytes()  # read it all at once so the file is never left locked
    except PermissionError:
        problems.append(Problem(
            None, None, None,
            f"The game is not allowed to read {name}. Another program may be using it.",
            f"Close {name} in Excel (and any other program that has it open), then try again."))
        return None
    except OSError as err:
        problems.append(Problem(
            None, None, None,
            f"{name} could not be read ({err.strerror or type(err).__name__}).",
            "Check that the file is not damaged and that the disk or USB stick is still connected."))
        return None
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # openpyxl warns about Excel extras it skips
            return load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        problems.append(Problem(
            None, None, None,
            f"{name} is not a valid Excel workbook. It may be damaged or saved in another format.",
            f"Open it in Excel and use File > Save As > Excel Workbook (*.xlsx), named {name}."))
        return None


def _read_content(workbook, problems: list[Problem]) -> Content:
    sheets = _find_sheets(workbook, problems)
    countries, codes = _read_countries(sheets["countries"])
    events, ids = _read_events(sheets["events"])
    results = _read_results(sheets["results"], countries, codes, events, ids)
    return Content(countries, events, results)


def _find_sheets(workbook, problems: list[Problem]) -> dict[str, _Sheet | None]:
    """Match sheet names ignoring case and spaces. A missing sheet becomes None."""
    by_name = {}
    for ws in workbook.worksheets:
        by_name.setdefault(ws.title.strip().lower(), ws)
    sheets = {}
    for name in HEADERS:
        ws = by_name.get(name)
        if ws is None:
            have = ", ".join(s.title for s in workbook.worksheets) or "none"
            problems.append(Problem(
                None, None, None,
                f"The sheet '{name}' is missing (sheets in the file: {have}).",
                f"Add a sheet named {name}, or rename the right sheet tab at the bottom of Excel to {name}."))
        sheets[name] = _Sheet(name, ws, problems) if ws is not None else None
    return sheets


class _Sheet:
    """One worksheet: finds the columns by their header and reads typed cells.

    Every bad cell adds a Problem that says where it is and how to fix it.
    """

    def __init__(self, name: str, ws, problems: list[Problem]):
        self.name = name
        self.problems = problems
        self.cols: dict[str, int] = {}       # header -> 0-based column index
        self.ok = self._find_columns(ws)     # False if a column is missing
        self.rows = list(self._data_rows(ws)) if self.ok else []

    def _find_columns(self, ws) -> bool:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        found = {}
        for index, value in enumerate(first):
            if value is not None:
                found.setdefault("_".join(str(value).lower().split()), index)
        wanted = HEADERS[self.name]
        missing = [h for h in wanted if h not in found]
        if len(missing) == len(wanted):
            names = ", ".join(wanted)
            self.add(1, None, f"Row 1 should hold the column names ({names}), but none were found.",
                     f"Type the column names into row 1, one per column: {names}.")
            return False
        for header in missing:
            self.add(1, None, f"The column '{header}' is missing.",
                     f"Type {header} into an empty cell of row 1 and fill in that column below it.")
        self.cols = {h: found[h] for h in wanted if h in found}
        return not missing

    def _data_rows(self, ws):
        """Yield (Excel row number, {header: value}), skipping rows that are empty."""
        width = max(self.cols.values()) + 1
        for number, cells in enumerate(ws.iter_rows(min_row=2, max_col=width, values_only=True), start=2):
            values = {h: cells[i] for h, i in self.cols.items()}
            if any(_as_text(v) for v in values.values()):
                yield number, values

    def column(self, header: str) -> str:
        return f"{get_column_letter(self.cols[header] + 1)} ({header})"

    def add(self, row: int | None, header: str | None, message: str, fix: str) -> None:
        column = self.column(header) if header else None
        self.problems.append(Problem(self.name, row, column, message, fix))

    def text(self, row: int, cells: dict, header: str) -> str | None:
        """The cell as tidy text, or None (with a problem) if it is empty."""
        value = cells[header]
        text = _as_text(value)
        if not text:
            message = "The cell only contains spaces." if _is_spaces(value) else "The cell is empty."
            self.add(row, header, message, f"Type {WHAT[header]} into it.")
            return None
        return text

    def whole_number(self, row: int, cells: dict, header: str) -> int | None:
        """The cell as an int, or None (with a problem) if it isn't a whole number."""
        value = cells[header]
        number = _as_int(value)
        if number is None:
            self.add(row, header, _number_problem(value), NUMBER_FIX)
        return number

    def count_problem(self, count: int, limits: tuple[int, int], singular: str, plural: str) -> None:
        low, high = limits
        if not low <= count <= high:
            word = singular if count == 1 else plural
            self.add(None, None, f"The sheet has {count} {word}, but the game needs {low} to {high} {plural}.",
                     f"Add or delete {singular} rows so there are between {low} and {high}.")


def _read_countries(sheet: _Sheet | None) -> tuple[list[Country], dict[str, int] | None]:
    """Return the valid countries and {code: row}, or None if the codes can't be trusted."""
    if sheet is None or not sheet.ok:
        return [], None
    countries, codes, complete = [], {}, True
    for row, cells in sheet.rows:
        code = sheet.text(row, cells, "code")
        name = sheet.text(row, cells, "name")
        color = _read_color(sheet, row, cells)
        card = _read_card(sheet, row, cells)
        if code is None:
            complete = False
            continue
        code = code.upper()
        if code in codes:
            sheet.add(row, "code", f"The code '{code}' is already used in row {codes[code]}.",
                      "Each country needs its own code. Change or delete one of the two rows.")
            continue
        codes[code] = row
        if None not in (name, color, card):
            countries.append(Country(code, name, color, card))
    sheet.count_problem(len(sheet.rows), COUNTRY_LIMITS, "country", "countries")
    return countries, codes if complete else None


def _read_color(sheet: _Sheet, row: int, cells: dict) -> tuple[int, int, int] | None:
    value = cells["color"]
    text = sheet.text(row, cells, "color")
    if text is None:
        return None
    if isinstance(value, int) and not isinstance(value, bool) and 0 <= value <= 999999:
        text = f"{value:06d}"  # Excel turns a color like 004080 typed without # into a number
    hex_digits = text.removeprefix("#")
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", hex_digits):
        sheet.add(row, "color", f"'{_short(text)}' is not a color code.",
                  "Type # and 6 hex digits (0-9, A-F), for example #B3202A.")
        return None
    return tuple(int(hex_digits[i:i + 2], 16) for i in (0, 2, 4))


def _read_card(sheet: _Sheet, row: int, cells: dict) -> tuple[str, ...] | None:
    text = sheet.text(row, cells, "card")
    if text is None:
        return None
    facts = tuple(part.strip() for part in text.split(";") if part.strip())
    low, high = FACT_LIMITS
    if not low <= len(facts) <= high:
        word = "fact" if len(facts) == 1 else "facts"
        sheet.add(row, "card", f"The card has {len(facts)} {word}, but it needs {low} or {high}.",
                  f"Write {low} or {high} short facts separated by semicolons (;), "
                  "for example: Suez Canal; Pyramids; Big wheat importer")
        return None
    return facts


def _read_events(sheet: _Sheet | None) -> tuple[list[Event], dict[str, int] | None]:
    """Return the valid events and {id: row}, or None if the ids can't be trusted."""
    if sheet is None or not sheet.ok:
        return [], None
    events, ids, complete = [], {}, True
    for row, cells in sheet.rows:
        event_id = sheet.text(row, cells, "id")
        title = sheet.text(row, cells, "title")
        kind = sheet.text(row, cells, "kind")
        description = sheet.text(row, cells, "description")
        options = tuple(sheet.text(row, cells, h) for h in OPTIONS)
        if kind is not None:
            kind = kind.lower()
            if kind not in ("good", "bad"):
                sheet.add(row, "kind", f"The kind '{_short(cells['kind'])}' is not good or bad.",
                          "Type good (a lucky event) or bad (a crisis).")
                kind = None
        if event_id is None:
            complete = False
            continue
        if event_id in ids:
            sheet.add(row, "id", f"The id '{event_id}' is already used in row {ids[event_id]}.",
                      "Each event needs its own id. Change one of them (and its rows on the results sheet).")
            continue
        ids[event_id] = row
        if None not in (title, kind, description, *options):
            events.append(Event(event_id, title, kind, description, options))
    sheet.count_problem(len(sheet.rows), EVENT_LIMITS, "event", "events")
    return events, ids if complete else None


def _read_results(sheet, countries, codes, events, ids) -> dict[tuple[str, str], Result]:
    """Read the results rows and check there is exactly one per (event, country).

    codes and ids are None when the other sheet had problems with its codes or ids;
    then the cross-checks are skipped so one mistake isn't reported many times.
    """
    if sheet is None or not sheet.ok:
        return {}
    if not sheet.rows:
        sheet.add(None, None, "The sheet has no rows.",
                  "Add one row per event and country: event_id, country, three option values and a reason.")
        return {}
    results, seen = {}, {}
    unsure_ids, unsure_codes, unsure_all = set(), set(), False  # rows whose key is broken
    for row, cells in sheet.rows:
        event_id = sheet.text(row, cells, "event_id")
        code = sheet.text(row, cells, "country")
        code = code.upper() if code else None
        if event_id and ids is not None and event_id not in ids:
            sheet.add(row, "event_id", f"There is no event with id '{_short(event_id)}' on the events sheet.",
                      f"Use one of the event ids from the events sheet: {', '.join(ids)}.")
        if code and codes is not None and code not in codes:
            sheet.add(row, "country", f"There is no country with code '{_short(code)}' on the countries sheet.",
                      f"Use one of the codes from the countries sheet: {', '.join(codes)}.")
        id_ok = ids is not None and event_id in ids
        code_ok = codes is not None and code in codes
        if not (id_ok and code_ok):
            if id_ok:
                unsure_ids.add(event_id)
            elif code_ok:
                unsure_codes.add(code)
            else:
                unsure_all = True
        elif (event_id, code) in seen:
            sheet.add(row, None, f"Event {event_id} and country {code} already have a row (row {seen[event_id, code]}).",
                      "Keep one row per event and country: delete this row or correct its event_id or country.")
            continue
        else:
            seen[event_id, code] = row
        values = tuple(sheet.whole_number(row, cells, h) for h in OPTIONS)
        reason = sheet.text(row, cells, "reason")
        if id_ok and code_ok and None not in values and reason is not None:
            results[event_id, code] = Result(values, reason)

    if ids is None or codes is None or unsure_all:
        return results
    titles = {e.id: f" ({e.title})" for e in events}
    names = {c.code: f" ({c.name})" for c in countries}
    for event_id in ids:
        for code in codes:
            if (event_id, code) in seen or event_id in unsure_ids or code in unsure_codes:
                continue
            sheet.add(None, None,
                      f"There is no row for event {event_id}{titles.get(event_id, '')} "
                      f"and country {code}{names.get(code, '')}.",
                      f"Add a row with event_id {event_id}, country {code}, the three option values and a reason.")
    return results


# --- Cell helpers -------------------------------------------------------------

def _as_text(value) -> str:
    """Cell value as tidy text: whole numbers without '.0', spaces and line breaks collapsed."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel may store the id 1 as 1.0
    return " ".join(str(value).split())


def _is_spaces(value) -> bool:
    return isinstance(value, str) and value != "" and not value.strip()


def _as_int(value) -> int | None:
    """Accept 2, 2.0, "2", "+2", "-1", " - 1", and "−1" or "–1" (minus sign, en dash)."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    if isinstance(value, str):
        text = value.strip().replace("−", "-").replace("–", "-")
        match = re.fullmatch(r"([+-]?)\s*([0-9]+)", text)
        if match:
            return int(match.group(1) + match.group(2))
    return None


def _number_problem(value) -> str:
    """Say why a value is not a whole number."""
    if value is None or value == "":
        return "The cell is empty."
    if _is_spaces(value):
        return "The cell only contains spaces."
    if isinstance(value, bool):
        return f"{_as_text(value)} is not a number."
    if isinstance(value, float):
        percent = round(value * 100, 6)
        if abs(value) < 1 and percent.is_integer():
            return f"{value:g} is not a whole number (Excel stores {int(percent)}% as {value:g})."
        return f"{value:g} is not a whole number."
    return f"'{_short(value)}' is not a whole number."


def _short(value, limit: int = 40) -> str:
    text = _as_text(value)
    return text if len(text) <= limit else text[:limit - 3] + "..."
